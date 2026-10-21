"""Step 3 — calculations and the Excel workbook.

Builds docs/dati/aizsardzibas_izdevumi.xlsx (Latvian sheet and column names) with
year-by-year tables, period averages, group averages and a simple correlation, plus
data/processed/kopsavilkums.json with the key numbers used by the web page.
"""
import json
import time

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from kopigais import (BALTICS, DOCS, GROUP_ORDER, PROCESSED, ROOT, SOURCE_ES, SOURCE_WB, WAR,
                      eu_weighted, latest_year, load, periods)


def wide(df, col, last, digits=None):
    t = df[(df["year"] >= 2000) & (df["year"] <= last)].pivot_table(index=["grupa", "valsts"], columns="year", values=col)
    t = t.reindex(sorted(t.index, key=lambda k: (GROUP_ORDER.index(k[0]) if k[0] in GROUP_ORDER else 99, k[1])))
    t.index.names = ["Reģions", "Valsts"]
    t.columns = [int(c) for c in t.columns]
    return t.round(digits) if digits is not None else t


def period_table(df, last):
    rows = []
    for (iso, name, grp), g in df.groupby(["iso3", "valsts", "grupa"]):
        g = g.set_index("year")
        r = {"Valsts": name, "Reģions": grp}
        for a, b, label in periods(last):
            r[f"Vid. % no IKP {label}"] = g.loc[a:b, "mil_gdp"].mean()
        r["% no IKP 2014"] = g["mil_gdp"].get(2014, np.nan)
        r[f"% no IKP {last}"] = g["mil_gdp"].get(last, np.nan)
        r[f"Izmaiņa 2014→{last} (proc. punkti)"] = r[f"% no IKP {last}"] - r["% no IKP 2014"]
        real14, reall = g["mil_real_usd_2015"].get(2014, np.nan), g["mil_real_usd_2015"].get(last, np.nan)
        r[f"Reālo izdevumu pieaugums 2014→{last}, %"] = (reall / real14 - 1) * 100 if real14 else np.nan
        r[f"Uz 1 iedz., USD {last}"] = g["mil_per_person_usd"].get(last, np.nan)
        r[f"≥ 2% no IKP {last}"] = "jā" if r[f"% no IKP {last}"] >= 2 else ("nē" if pd.notna(r[f"% no IKP {last}"]) else "–")
        r["ES dalībvalsts"] = "jā" if g["es"].iloc[0] == 1 else "nē"
        r[f"NATO dalībvalsts {last}"] = "jā" if g["nato"].get(last, 0) == 1 else "nē"
        rows.append(r)
    t = pd.DataFrame(rows)
    t["_o"] = t["Reģions"].map(lambda x: GROUP_ORDER.index(x) if x in GROUP_ORDER else 99)
    return t.sort_values(["_o", "Valsts"]).drop(columns="_o").round(2)


def group_table(df, last):
    d = df[~df["iso3"].isin(WAR) & (df["year"] <= last)]
    t = d.pivot_table(index="year", columns="grupa", values="mil_gdp", aggfunc="mean")
    t = t[[g for g in GROUP_ORDER if g in t.columns]]
    t["ES-27 (svērtais)"] = eu_weighted(df)
    t["ES-27 (vienkāršais vid.)"] = df[df["es"] == 1].groupby("year")["mil_gdp"].mean()
    t.index.name = "Gads"
    return t.loc[2000:last].round(2)


def correlation_table(df, last):
    d = df[~df["iso3"].isin(WAR) & df["year"].between(2000, last)]
    g = d.groupby(["valsts", "grupa"]).agg(mil=("mil_gdp", "mean"), growth=("gdp_growth", "mean")).dropna().reset_index()
    r = g["mil"].corr(g["growth"])
    g.columns = ["Valsts", "Reģions", f"Vid. aizsardzības izdevumi, % no IKP (2000–{last})", f"Vid. reālā IKP izaugsme, % gadā (2000–{last})"]
    return g.round(2), r


def style_sheet(ws, title, note, first_data_row):
    ws.insert_rows(1, 3)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = note
    ws["A2"].font = Font(italic=True, color="666666")
    hdr = first_data_row + 3
    for cell in ws[hdr]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E6ECF8")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, col in enumerate(ws.iter_cols(min_row=hdr, max_row=ws.max_row), start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:60])
        ws.column_dimensions[get_column_letter(i)].width = min(max(9, width + 2), 34)
    ws.freeze_panes = ws.cell(row=hdr + 1, column=3)


def clean(o):
    """Make numbers JSON-safe (NaN -> null, numpy -> plain Python)."""
    if isinstance(o, dict):
        return {k: clean(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [clean(v) for v in o]
    if isinstance(o, (np.floating, float)):
        return None if np.isnan(o) else float(o)
    if isinstance(o, np.integer):
        return int(o)
    return o


def main():
    df = load()
    last = latest_year(df)
    dl_date = (ROOT / "data" / "raw" / "lejupielades_datums.txt").read_text().strip() if (ROOT / "data" / "raw" / "lejupielades_datums.txt").exists() else time.strftime("%Y-%m-%d")
    out = DOCS / "dati"
    out.mkdir(parents=True, exist_ok=True)

    sheets = [
        ("% no IKP", wide(df, "mil_gdp", last, 2), "Aizsardzības (militārie) izdevumi, % no IKP", SOURCE_WB),
        ("Uz 1 iedzīvotāju (USD)", wide(df, "mil_per_person_usd", last, 0), "Aizsardzības izdevumi uz vienu iedzīvotāju, ASV dolāri faktiskajās cenās", SOURCE_WB),
        ("Reālie izdevumi (milj. USD)", wide(df.assign(x=df["mil_real_usd_2015"] / 1e6), "x", last, 1), "Aizsardzības izdevumi salīdzināmajās cenās, milj. ASV dolāru (2015. gada cenās)", SOURCE_WB + "; reālie izdevumi = % no IKP × IKP 2015. g. cenās"),
        ("% no valdības izdevumiem", wide(df, "mil_gov", last, 2), "Aizsardzības izdevumi, % no vispārējās valdības izdevumiem", SOURCE_WB),
        ("Eurostat % no IKP", wide(df, "es_cofog_gdp", last + 1, 2), "Valdības izdevumi aizsardzībai pēc funkcijām (COFOG GF02), % no IKP", SOURCE_ES),
    ]
    per = period_table(df, last)
    grp = group_table(df, last)
    corr, r = correlation_table(df, last)

    readme = pd.DataFrame({"Lasi mani": [
        "Aizsardzības izdevumi Eiropā, 2000–" + str(last),
        f"Dati lejupielādēti: {dl_date}. Jaunākais gads ar pilnīgiem datiem: {last}.",
        "",
        "Lapas:",
        "% no IKP — militārie izdevumi kā daļa no iekšzemes kopprodukta (galvenais rādītājs; NATO mērķis ir 2 %).",
        "Uz 1 iedzīvotāju (USD) — izdevumi ASV dolāros, dalīti ar iedzīvotāju skaitu (faktiskajās cenās).",
        "Reālie izdevumi — izdevumi 2015. gada cenās, lai salīdzinātu apjomu laika gaitā bez inflācijas ietekmes.",
        "% no valdības izdevumiem — cik liela daļa no valsts budžeta izdevumiem aiziet aizsardzībai.",
        "Eurostat % no IKP — tas pats no ES valdības finanšu statistikas (COFOG). Definīcija atšķiras no SIPRI/NATO, tāpēc skaitļi nedaudz atšķiras.",
        "Periodu vidējie — vidējie rādītāji 2000–2013, 2014–2021 un no 2022, izmaiņas un 2 % mērķa izpilde.",
        "Reģionu vidējie — reģionu vienkāršie vidējie un ES-27 svērtais vidējais (kopējie izdevumi / kopējais IKP).",
        "Korelācija — vidējie izdevumi pret vidējo IKP izaugsmi; tikai aprakstoša sakarība, nevis cēloņsakarība.",
        "",
        "Ukraina, Krievija un Baltkrievija ir tabulās, bet nav iekļautas reģionu vidējos un korelācijā (kara ietekme).",
        "Avoti: Pasaules Banka (militāro izdevumu dati no SIPRI), https://data.worldbank.org; Eurostat gov_10a_exp, https://ec.europa.eu/eurostat.",
    ]})

    path = out / "aizsardzibas_izdevumi.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        readme.to_excel(xw, sheet_name="Lasi mani", index=False)
        xw.sheets["Lasi mani"].column_dimensions["A"].width = 120
        for name, t, title, note in sheets:
            t.to_excel(xw, sheet_name=name)
            style_sheet(xw.sheets[name], title, note, 1)
        per.to_excel(xw, sheet_name="Periodu vidējie", index=False)
        style_sheet(xw.sheets["Periodu vidējie"], f"Periodu vidējie un izmaiņas, 2000–{last}", SOURCE_WB, 1)
        grp.to_excel(xw, sheet_name="Reģionu vidējie")
        style_sheet(xw.sheets["Reģionu vidējie"], "Reģionu vidējie aizsardzības izdevumi, % no IKP", SOURCE_WB + " (bez Ukrainas, Krievijas, Baltkrievijas)", 1)
        corr.to_excel(xw, sheet_name="Korelācija", index=False)
        style_sheet(xw.sheets["Korelācija"], f"Aizsardzības izdevumi un IKP izaugsme (Pīrsona r = {r:.2f})".replace(".", ","), "Aprakstoša sakarība starp valstīm; nenozīmē cēloņsakarību. " + SOURCE_WB, 1)

    df[df["year"] <= last + 1].to_csv(out / "dati_visi.csv", index=False)

    # key numbers for the web page
    piv = df.pivot_table(index="iso3", columns="year", values="mil_gdp")
    names = df.drop_duplicates("iso3").set_index("iso3")["valsts"]
    eu = eu_weighted(df)
    nato_eu = df[(df["year"] == last) & (df["nato"] == 1)]
    nato14 = df[(df["year"] == 2014) & (df["iso3"].isin(nato_eu["iso3"]))]
    rank = piv[last].drop(list(WAR), errors="ignore").dropna().sort_values(ascending=False)
    facts = {
        "latest": last, "downloaded": dl_date,
        "baltics": {iso: {"name": names[iso], "y2000": piv.get(2000, {}).get(iso), "y2014": piv[2014].get(iso), "last": piv[last].get(iso),
                           "rank": int(list(rank.index).index(iso) + 1) if iso in rank.index else None} for iso in BALTICS},
        "n_ranked": int(len(rank)),
        "top3": [[names[i], float(v)] for i, v in rank.head(3).items()],
        "eu_2014": float(eu.get(2014, np.nan)), "eu_last": float(eu.get(last, np.nan)),
        "nato_eur_over2_last": int((nato_eu["mil_gdp"] >= 2).sum()), "nato_eur_n_last": int(nato_eu["mil_gdp"].notna().sum()),
        "nato_eur_over2_2014": int((nato14["mil_gdp"] >= 2).sum()),
        "corr_r": float(r),
    }
    (PROCESSED / "kopsavilkums.json").write_text(json.dumps(clean(facts), ensure_ascii=False, indent=1))
    print(f"Excel workbook saved to {path.relative_to(ROOT)} (latest year {last})")


if __name__ == "__main__":
    main()
